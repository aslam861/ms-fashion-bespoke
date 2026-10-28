import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def add_paginated_code_file(doc, filepath, display_name, lines_per_page=36, max_total_lines=None):
    if not os.path.exists(filepath):
        print(f"Warning: {filepath} not found")
        return 0

    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        all_lines = [line.rstrip('\r\n') for line in f.readlines()]

    if max_total_lines and len(all_lines) > max_total_lines:
        all_lines = all_lines[:max_total_lines]

    total_chunks = (len(all_lines) + lines_per_page - 1) // lines_per_page
    print(f"Adding {display_name}: {len(all_lines)} lines across {total_chunks} pages")

    for c_idx in range(total_chunks):
        chunk = all_lines[c_idx * lines_per_page : (c_idx + 1) * lines_per_page]
        
        # New page for each code chunk to ensure strict pagination like Shifa's report
        doc.add_page_break()

        # Filename title on the page
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(4)
        run_title = p_title.add_run(f"{display_name} (Page {c_idx+1} of {total_chunks})")
        run_title.bold = True
        run_title.font.name = 'Times New Roman'
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(0x5C, 0x1A, 0x21)

        # Single bordered box for the page
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="8" w:space="0" w:color="5C1A21"/>'
            f'<w:left w:val="single" w:sz="8" w:space="0" w:color="5C1A21"/>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="5C1A21"/>'
            f'<w:right w:val="single" w:sz="8" w:space="0" w:color="5C1A21"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAF8F5"/>')
        tcPr.append(shd)

        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        
        chunk_text = "\n".join(chunk)
        r = p.add_run(chunk_text)
        r.font.name = 'Courier New'
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x1F, 0x20, 0x25)

    return total_chunks

print("paginated code helper ready")
