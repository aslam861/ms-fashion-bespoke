import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from docx_helpers import set_cell_background, set_cell_margins, add_page_border, add_page_number_to_section

def build_report():
    doc = docx.Document()

    # Base Section Setup
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25) # 1.25 for binding
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.27)  # A4
    section.page_height = Inches(11.69)
    add_page_border(section)
    add_page_number_to_section(section)

    # Style definitions
    PRIMARY_COLOR = RGBColor(0x5C, 0x1A, 0x21) # Maroon
    GOLD_COLOR = RGBColor(0xA9, 0x81, 0x2F)    # Gold
    DARK_TEXT = RGBColor(0x2B, 0x21, 0x1D)     # Charcoal body

    # Helper functions for text
    def add_title(text, font_size=16, color=PRIMARY_COLOR, space_before=12, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        return p

    def add_h1(text):
        doc.add_page_break()
        return add_title(text, font_size=16, color=PRIMARY_COLOR, space_before=18, space_after=14)

    def add_h2(text):
        return add_title(text, font_size=13, color=PRIMARY_COLOR, space_before=14, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.italic = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = DARK_TEXT
        return p

    def add_p(text, bold_prefix='', italic=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(12)
            r_bold.font.color.rgb = DARK_TEXT
        run = p.add_run(text)
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = DARK_TEXT
        return p

    def add_bullet(text, bold_prefix=''):
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(12)
            r_bold.font.color.rgb = DARK_TEXT
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = DARK_TEXT
        return p

    def add_image_figure(img_path, caption_text, width=Inches(5.8)):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(img_path, width=width)

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            run = p_cap.add_run(caption_text)
            run.bold = True
            run.italic = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.color.rgb = PRIMARY_COLOR
        else:
            add_p(f"[{caption_text}]")

    def add_code_box(code_text, filename=''):
        if filename:
            p_fn = doc.add_paragraph()
            p_fn.paragraph_format.space_before = Pt(8)
            p_fn.paragraph_format.space_after = Pt(2)
            run_fn = p_fn.add_run(f"Source: {filename}")
            run_fn.bold = True
            run_fn.font.name = 'Times New Roman'
            run_fn.font.size = Pt(11)
            run_fn.font.color.rgb = PRIMARY_COLOR

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F4F4F6")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        # Add code paragraph inside cell
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(code_text)
        run.font.name = 'Courier New'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x25)
        
        # empty space after table
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(2)
        p_sp.paragraph_format.space_after = Pt(6)

    print("Setup completed, now writing chapters...")
    return doc, add_title, add_h1, add_h2, add_h3, add_p, add_bullet, add_image_figure, add_code_box

if __name__ == '__main__':
    doc, add_title, add_h1, add_h2, add_h3, add_p, add_bullet, add_image_figure, add_code_box = build_report()
    doc.save('test_report_init.docx')
    print('Initial test document saved successfully')
