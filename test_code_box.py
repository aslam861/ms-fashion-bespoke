import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
section = doc.sections[0]
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.25)
section.right_margin = Inches(1.0)

def add_full_code_file(doc, filepath, display_name):
    # Header for code file
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(14)
    p_head.paragraph_format.space_after = Pt(4)
    r_head = p_head.add_run(display_name)
    r_head.bold = True
    r_head.font.name = 'Times New Roman'
    r_head.font.size = Pt(11)
    r_head.font.color.rgb = RGBColor(0x5C, 0x1A, 0x21)

    # Read lines
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        lines = [line.rstrip('\r\n') for line in f.readlines()]

    # To ensure clean page-by-page flow like Shifa's report,
    # we chunk lines into ~38-40 lines per box/page
    chunk_size = 40
    for i in range(0, len(lines), chunk_size):
        chunk = lines[i:i+chunk_size]
        tbl = doc.add_table(rows=1, cols=1)
        cell = tbl.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="5C1A21"/>'
            f'<w:left w:val="single" w:sz="6" w:space="0" w:color="5C1A21"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="5C1A21"/>'
            f'<w:right w:val="single" w:sz="6" w:space="0" w:color="5C1A21"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAF8F5"/>')
        tcPr.append(shd)

        # First paragraph in cell
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(0)
        
        # Combine chunk text
        chunk_text = "\n".join(chunk)
        r = p.add_run(chunk_text)
        r.font.name = 'Courier New'
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x20, 0x20, 0x25)

        # Spacing after table
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(2)
        p_sp.paragraph_format.space_after = Pt(8)

add_full_code_file(doc, 'frontend/src/components/AtelierChatbot.jsx', 'AtelierChatbot.jsx')
doc.save('test_chunked_code.docx')
print('test_chunked_code.docx generated successfully!')
